"""
TMS-RAT report — built-in second-level (group-level) MEP-CMAP add-on.

Reads the per-recording analysis records that every first-level run writes
(``results/report/<prefix>_reporting.json``) and assesses the STUDY against the
TMS-RAT v1.0 reporting checklist, emitting:

    <study>_tms-rat.csv       the OUTPUT sheet format, one row per item
    <study>_tms-rat.xlsx      the official workbook with INPUT filled in
    <study>_methods.md        a methods draft with a marker at every gap

Why group level: TMS-RAT assesses one protocol within one study, and its
quantitative items are rated Full only when a central tendency AND a dispersion
are reported. A single recording cannot supply a dispersion across
participants, so a per-file assessment would rate those items Part for every
study however completely it was reported.

The rules live in ``mep_cmap.tms_rat``; this file is the I/O around them, so
the assessment can be tested without a study on disk.
"""

import csv
import glob
import json
import os
import shutil

# Absolute, not relative. Add-ons are loaded with spec_from_file_location under
# a synthetic module name and therefore have NO parent package, so a relative
# import raises at import time and the add-on is silently skipped as broken.
try:
    from mep_cmap import tms_rat as TR
    from mep_cmap.reporting import SIDECAR_SUFFIX, load_tms_rat_mapping
    _IMPORT_ERROR = None
except Exception as _ex:                  # pragma: no cover - packaging failure
    TR = None
    SIDECAR_SUFFIX = "reporting.json"
    load_tms_rat_mapping = lambda *_a, **_k: None       # noqa: E731
    _IMPORT_ERROR = _ex

ADDON_NAME        = "tms_rat_report"
ADDON_DESCRIPTION = ("Assess the study against the TMS-RAT v1.0 reporting "
                     "checklist and draft a methods section, from the "
                     "per-recording analysis records")
ADDON_VERSION     = "1.0.0"
ADDON_AUTHOR      = "MEP-CMAP Analyser (built-in)"
ADDON_SCOPE       = "group_level"

ADDON_SETTINGS = [
    {"key": "tms_rat_study", "label": "Study name or number",
     "help": "Written into the Study column and the workbook's INPUT!C3.",
     "type": "str", "default": ""},
    {"key": "tms_rat_id", "label": "PMID, DOI or URL",
     "help": ("Identifies the paper being assessed. Written to the ID "
              "column and the workbook's INPUT!C2."),
     "type": "str", "default": ""},
    {"key": "tms_rat_rater", "label": "Rater initials",
     "help": ("Who is responsible for the ratings. The tool fills in what it "
              "can derive; a human still has to check it, and the sheet has "
              "to say who."),
     "type": "str", "default": ""},
    {"key": "tms_rat_protocol", "label": "Protocol assessed",
     "help": ("Which TMS protocol this study used. It decides which sections "
              "the checklist requires. Left blank, the Required column is "
              "blank too rather than assuming every section applies."),
     "type": "str", "default": "Motor-evoked potentials",
     "choices": ["", "Single pulse TMS", "Motor-evoked potentials", "Motor",
                 "Afferent conditioning", "Paired-associate conditioning",
                 "Paired-pulse TMS", "Repetitive TMS"]},
]

#: Where the official workbook ships, so the filled copy is the real thing
#: rather than a reconstruction of it. Absent in a source checkout that has not
#: fetched it; the add-on then writes the CSV and methods and says so.
_WORKBOOK = os.path.join(os.path.dirname(__file__),
                         "..", "..", "..", "tools",
                         "TMS_Reporting_Assessment_Tool_v1.0.xlsx")


def find_records(root, log=None):
    """Every per-recording analysis record under a derivatives tree.

    Searched recursively rather than in one known folder: the records sit in
    each recording's own ``results/report/``, and a study analysed before the
    family layout existed keeps them flat beside its other results.
    """
    paths = sorted(glob.glob(os.path.join(root, "**", f"*_{SIDECAR_SUFFIX}"),
                             recursive=True))
    records = []
    for p in paths:
        try:
            with open(p, encoding="utf-8") as fh:
                rec = json.load(fh)
        except Exception as exc:                # noqa: BLE001
            if log:
                log(f"{ADDON_NAME}: could not read {os.path.basename(p)} "
                    f"({type(exc).__name__}: {exc}); skipped.")
            continue
        if rec.get("schema") == "mep-cmap-reporting":
            rec["_path"] = p
            records.append(rec)
    return records


def unrecorded_recordings(root, records):
    """Recordings with results but no analysis record.

    A study-level assessment silently based on one of three recordings is the
    worst failure this add-on can have: its whole justification is that a
    dispersion needs several recordings, and it would report one file's numbers
    as the study's without a word. Anything analysed before the record existed,
    or by an older install, has results and no record -- so the gap is normal,
    invisible, and has to be said out loud.
    """
    have = set()
    for r in records:
        stem = os.path.basename(r.get("_path", ""))
        if stem.endswith("_" + SIDECAR_SUFFIX):
            have.add(stem[:-(len(SIDECAR_SUFFIX) + 1)])
    seen = set()
    for p in glob.glob(os.path.join(root, "**", "*_trials.csv"),
                       recursive=True):
        name = os.path.basename(p)
        if name.endswith("_trials_selected.csv"):
            continue
        seen.add(name[:-len("_trials.csv")])
    return sorted(seen - have)


def find_participants(root):
    """The study's participants.tsv, if it has one.

    Searched in the places a NIBS-BIDS study actually puts it, nearest first.
    The analysis writes into ``derivatives/``, and BIDS puts participants.tsv
    at the root of the raw dataset -- which, when the tool has BIDS-ified the
    source itself, is ``rawdata/`` beside ``derivatives/`` rather than the
    study folder. Missing that one reports "no participants.tsv in the
    dataset" for a study that has one, which is a false statement about the
    study rather than a gap in it.
    """
    parent = os.path.dirname(os.path.normpath(root))
    candidates = [
        os.path.join(root, "participants.tsv"),
        os.path.join(parent, "participants.tsv"),
        os.path.join(parent, "rawdata", "participants.tsv"),
        os.path.join(parent, "sourcedata", "participants.tsv"),
    ]
    for candidate in candidates:
        if os.path.isfile(candidate):
            try:
                with open(candidate, encoding="utf-8") as fh:
                    return TR.parse_participants_tsv(fh.read()), candidate
            except Exception:                   # noqa: BLE001
                return [], candidate
    return [], None


def fill_workbook(results, out_path, *, study="", study_id="", rater="",
                  protocol="", source=None):
    """Copy the official workbook and fill INPUT with what was derived.

    INPUT is filled, not OUTPUT. OUTPUT is entirely formulas that read INPUT
    and recalculate when the file is opened, so writing values into it would
    replace the checklist's own logic with ours and the two could then
    disagree without anyone noticing.

    Returns the path, or None when the workbook is not available. openpyxl and
    the workbook are both optional: a study still gets its CSV and its methods
    draft without them.
    """
    src = source or _WORKBOOK
    if not os.path.isfile(src):
        return None
    try:
        import openpyxl
    except ImportError:
        return None

    shutil.copyfile(src, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = wb["INPUT"]
    if study_id:
        ws["C2"] = study_id
    if study:
        ws["C3"] = study
    if rater:
        ws["C4"] = rater
    if protocol:
        ws["C5"] = protocol

    for r in results:
        row = r.get("input_row")
        if not row:
            continue
        # A rating the tool did not derive is left EMPTY rather than written as
        # "Missing". An empty cell is an unanswered question the rater still
        # sees; "Missing" is an answer, and one the tool is not entitled to
        # give on an item it never looked at.
        if r["rating"] not in (TR.BLANK, ""):
            ws.cell(row=row, column=3).value = r["rating"]
        if r["comment"] not in (TR.BLANK, ""):
            ws.cell(row=row, column=4).value = r["comment"]
        ws.cell(row=row, column=10).value = r["timestamp"]

    wb.save(out_path)
    return out_path


def run(context):
    log = context.log
    cfg = context.config or {}
    root = context.results_dir

    if TR is None:
        log(f"{ADDON_NAME}: could not import mep_cmap.tms_rat "
            f"({_IMPORT_ERROR}), so no assessment can be made.")
        return []

    mapping = load_tms_rat_mapping()
    if mapping is None:
        log(f"{ADDON_NAME}: the TMS-RAT item mapping is missing from the "
            f"install, so no assessment can be made.")
        return []

    records = find_records(root, log=log)
    if not records:
        log(f"{ADDON_NAME}: no analysis records found under {root}. "
            f"Re-run First Level to write them.")
        return []

    participants, ppath = find_participants(root)
    if ppath is None:
        log(f"{ADDON_NAME}: no participants.tsv found. The participant items "
            f"will be Missing: this tool does not collect participant data.")
    else:
        log(f"{ADDON_NAME}: read {len(participants)} participant row(s) from "
            f"{ppath}")

    # Said before anything is written, and repeated into the methods draft,
    # because an assessment covering a third of the study looks exactly like
    # one covering all of it.
    orphans = unrecorded_recordings(root, records)
    coverage_note = ""
    if orphans:
        shown = ", ".join(orphans[:5])
        more = f" (+{len(orphans) - 5} more)" if len(orphans) > 5 else ""
        coverage_note = (
            f"{len(records)} of {len(records) + len(orphans)} recordings in "
            f"this tree carry an analysis record. This assessment describes "
            f"only those {len(records)}. Without a record from the others "
            f"— {shown}{more} — re-run First Level on them so the study is "
            f"assessed as a whole rather than as a sample of itself.")
        log(f"{ADDON_NAME}: ⚠️  {coverage_note}")

    study = str(cfg.get("tms_rat_study") or "").strip()
    protocol = str(cfg.get("tms_rat_protocol") or "").strip()
    stem = study or "study"
    out_dir = os.path.join(root, "report")
    os.makedirs(out_dir, exist_ok=True)

    try:
        results, faults = TR.assess(
            records, mapping, participants=participants, study=study,
            rater=str(cfg.get("tms_rat_rater") or "").strip(),
            study_id=str(cfg.get("tms_rat_id") or "").strip(),
            protocol=protocol or None)
    except TR.AnalysisFault as exc:
        # Refused, not worked around. Pooling these would produce numbers
        # describing no analysis that was ever run.
        log(f"{ADDON_NAME}: REFUSED to assess these recordings together. {exc}")
        return []

    if coverage_note:
        faults.insert(0, coverage_note)

    written = []

    csv_path = os.path.join(out_dir, f"{stem}_tms-rat.csv")
    with open(csv_path, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=TR.OUTPUT_COLUMNS)
        w.writeheader()
        w.writerows(TR.output_rows(results))
    written.append(csv_path)

    md_path = os.path.join(out_dir, f"{stem}_methods.md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write(TR.methods_markdown(
            results, faults, study=study, protocol=protocol,
            tool_version=records[0].get("tool_version") or "",
            detection_version=records[0].get("detection_version") or ""))
    written.append(md_path)

    xlsx_path = fill_workbook(
        results, os.path.join(out_dir, f"{stem}_tms-rat.xlsx"),
        study=study, study_id=str(cfg.get("tms_rat_id") or "").strip(),
        rater=str(cfg.get("tms_rat_rater") or "").strip(), protocol=protocol)
    if xlsx_path:
        written.append(xlsx_path)
    else:
        log(f"{ADDON_NAME}: the official workbook or openpyxl is not "
            f"available, so only the CSV and methods draft were written.")

    counts = TR.coverage(results)
    log(f"{ADDON_NAME}: {len(records)} recording(s) assessed against "
        f"{len(results)} items — "
        + ", ".join(f"{k} {v}" for k, v in sorted(counts.items()))
        + f" → {os.path.basename(csv_path)}")
    for f in faults:
        log(f"{ADDON_NAME}: ⚠️  {f}")
    return written
